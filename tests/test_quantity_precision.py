from uuid import uuid4
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from app.core.quantity import decimal_json, quantity_text
from app.models import Product, StockBalance, StockTransaction, InventoryMovement
from app.schemas.stock_schema import StockIn, StockOut, StockAdjust
from app.schemas.stock_balance_schema import StockBalanceCreate, StockBalanceAdjust
from app.schemas.inventory_transfer_schema import InventoryTransferItemCreate
from app.schemas.sales_return_schema import SalesReturnItemCreate
from app.repositories.dashboard_repository import DashboardRepository
from app.services.dashboard_service import get_stock_summary_service
from tests.test_stock import _create_product, _create_batch
from tests import test_purchase_order as purchase
from tests import test_sales_order as sales


CONTRACTS = [
    (StockIn, {"product_id": 1}, "quantity"),
    (StockOut, {"product_id": 1}, "quantity"),
    (StockAdjust, {"product_id": 1, "remark": "count"}, "new_quantity"),
    (StockBalanceCreate, {"product_id": 1, "warehouse_id": 1, "location_id": 1}, "on_hand_qty"),
    (StockBalanceCreate, {"product_id": 1, "warehouse_id": 1, "location_id": 1}, "reserved_qty"),
    (StockBalanceAdjust, {}, "on_hand_qty"),
    (StockBalanceAdjust, {}, "reserved_qty"),
    (InventoryTransferItemCreate, {"product_id": 1, "from_location_id": 1, "to_location_id": 2}, "quantity"),
    (SalesReturnItemCreate, {"product_id": 1}, "quantity"),
]


@pytest.mark.parametrize("model,payload,field", CONTRACTS)
@pytest.mark.parametrize("quantity", ["0.0001", "1000000000000000", "NaN", "Infinity", "-Infinity"])
def test_quantity_contract_rejects_unrepresentable(model, payload, field, quantity):
    with pytest.raises(ValidationError):
        model.model_validate({**payload, field: quantity})


@pytest.mark.parametrize("model,payload,field", CONTRACTS)
@pytest.mark.parametrize("quantity", ["0.001", "1.125", "999999999999999.999"])
def test_quantity_contract_preserves_exact_decimal(model, payload, field, quantity):
    parsed = model.model_validate({**payload, field: quantity})
    assert getattr(parsed, field) == Decimal(quantity)


def test_compatible_numeric_json_and_exact_large_fallback():
    assert decimal_json(Decimal("1.125")) == 1.125
    assert isinstance(decimal_json(Decimal("1.125")), float)
    assert decimal_json(Decimal("12.000")) == 12
    assert decimal_json(Decimal("999999999999999.999")) == "999999999999999.999"
    assert quantity_text(Decimal("0.001")) == "0.001"
    with pytest.raises(ValueError):
        quantity_text(Decimal("1.0001"))


def test_fractional_dashboard_and_report_export(client, admin_headers, db_session):
    product = _create_product(client, admin_headers, initial_stock="1.125")
    rows = get_stock_summary_service(DashboardRepository(db_session))
    row = next(r for r in rows if r["product_id"] == product["id"])
    assert row["stock_qty"] == Decimal("1.125")
    response = client.get("/api/v1/dashboard/stock-summary", headers=admin_headers)
    row = next(r for r in response.json()["data"]["items"] if r["product_id"] == product["id"])
    assert row["stock_qty"] == 1.125
    assert isinstance(row["stock_qty"], float)
    response = client.get("/api/v1/reports/export/stock", headers=admin_headers)
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content))
    row = next(r for r in workbook.active.iter_rows(min_row=2) if r[0].value == product["id"])
    assert row[3].value == "1.125"
    assert row[3].data_type == "s"


def test_boundary_quantity_report_and_excel_are_exact(client, admin_headers):
    product = _create_product(client, admin_headers, initial_stock="999999999999999.999")
    response = client.get("/api/v1/reports/stock-balance", headers=admin_headers)
    row = next(r for r in response.json() if r["product_id"] == product["id"])
    assert row["stock_qty"] == "999999999999999.999"
    chart = client.get("/api/v1/reports/chart/stock", headers=admin_headers)
    assert chart.status_code == 200
    quantities = [Decimal(str(row["stock_qty"])) for row in chart.json()]
    assert Decimal("999999999999999.999") in quantities
    assert quantities == sorted(quantities, reverse=True)
    response = client.get("/api/v1/reports/export/stock", headers=admin_headers)
    workbook = load_workbook(BytesIO(response.content))
    row = next(r for r in workbook.active.iter_rows(min_row=2) if r[0].value == product["id"])
    assert row[3].value == "999999999999999.999"


def test_excess_precision_request_leaves_inventory_untouched(client, admin_headers, db_session):
    product = _create_product(client, admin_headers, initial_stock="1.125")
    for endpoint, payload in [
        ("in", {"quantity": "0.0005"}),
        ("out-fefo", {"quantity": "0.0005"}),
        ("adjust", {"new_quantity": "0.0005", "remark": "count"}),
    ]:
        response = client.post(f"/api/v1/stock/{endpoint}", headers=admin_headers,
                               json={"product_id": product["id"], **payload})
        assert response.status_code == 422
    db_session.expire_all()
    assert db_session.get(Product, product["id"]).stock_qty == Decimal("1.125")
    assert db_session.query(StockTransaction).filter_by(product_id=product["id"]).count() == 1
    assert db_session.query(InventoryMovement).filter_by(product_id=product["id"]).count() == 1


def test_two_products_can_receive_same_lot_in_one_po(client, admin_headers):
    supplier = purchase._create_supplier(client, admin_headers)
    products = [purchase._create_product(client, admin_headers) for _ in range(2)]
    response = client.post("/api/v1/purchase-orders", headers=admin_headers, json={
        "supplier_id": supplier["id"], "items": [
            {"product_id": p["id"], "quantity": "1.125", "unit_price": 1} for p in products
        ],
    })
    assert response.status_code == 201
    order_id = response.json()["data"]["id"]
    purchase._confirm_purchase_order(client, admin_headers, order_id)
    items = []
    for product in products:
        item = purchase._receive_payload(product["id"], quantity="1.125", lot_no="SHARED-LOT")["items"][0]
        items.append(item)
    response = client.post(f"/api/v1/purchase-orders/{order_id}/receive", headers={**admin_headers, "Idempotency-Key": uuid4().hex}, json={"items": items})
    assert response.status_code == 200
    batches = response.json()["data"]["received_batches"]
    assert len({b["batch_id"] for b in batches}) == 2
    assert {b["lot_no"] for b in batches} == {"SHARED-LOT"}


def test_dashboard_repository_keeps_fractional_sums(client, admin_headers, db_session):
    repository = DashboardRepository(db_session)
    initial = repository.get_total_stock()
    initial_in = repository.get_transaction_quantity_sum({"IN"})
    _create_product(client, admin_headers, initial_stock="0.001")
    _create_product(client, admin_headers, initial_stock="1.125")
    assert repository.get_total_stock() - initial == Decimal("1.126")
    assert repository.get_transaction_quantity_sum({"IN"}) - initial_in == Decimal("1.126")


def test_batch_lot_scope_preserves_case_and_duplicate_rejection(client, admin_headers):
    first = _create_product(client, admin_headers)
    second = _create_product(client, admin_headers)
    lot = "SCOPE-" + uuid4().hex
    payload = {"quantity": "0.001", "lot_no": lot,
               "mfg_date": date.today().isoformat(),
               "expiry_date": (date.today() + timedelta(days=90)).isoformat()}
    for product in [first, second]:
        response = client.post("/api/v1/batches", headers=admin_headers,
                               json={**payload, "product_id": product["id"]})
        assert response.status_code == 201
    response = client.post("/api/v1/batches", headers=admin_headers,
                           json={**payload, "product_id": first["id"]})
    assert response.status_code == 409
    response = client.post("/api/v1/batches", headers=admin_headers,
                           json={**payload, "product_id": first["id"], "lot_no": lot.lower()})
    assert response.status_code == 201


def test_fractional_invoice_keeps_quantity_text(client, admin_headers, monkeypatch):
    product = sales._create_product(client, admin_headers)
    customer = sales._create_customer(client, admin_headers)
    sales._create_batch(client, admin_headers, product["id"], quantity="1.125", expiry_days=90)
    order = sales._create_sales_order(client, admin_headers, customer["id"], product["id"],
                                     quantity="1.125", unit_price=1)
    from app.routers import sales_order_router
    original_table = sales_order_router.Table
    captured = []

    def record_table(data, *args, **kwargs):
        captured.extend(data)
        return original_table(data, *args, **kwargs)

    monkeypatch.setattr(sales_order_router, "Table", record_table)
    response = client.get(f"/api/v1/sales-orders/{order['sales_order_id']}/invoice", headers=admin_headers)
    assert response.status_code == 200
    assert captured[1][1] == "1.125"
    assert response.content.startswith(b"%PDF")
