from datetime import date, timedelta
from pathlib import Path
from uuid import uuid4

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models import ProductBatch


def _create_product(
    client: TestClient,
    admin_headers: dict[str, str],
    *,
    stock_qty: int = 0,
    price: float = 100.00,
) -> dict:
    unique_value = uuid4().hex[:10].upper()

    response = client.post(
        "/api/v1/products",
        headers=admin_headers,
        json={
            "sku": f"REPORT-{unique_value}",
            "barcode": f"881{unique_value}",
            "product_name": (
                f"Report Product {unique_value}"
            ),
            "price": price,
            "stock_qty": stock_qty,
            "category_id": None,
            "track_batch": True,
            "track_expiry": True,
        },
    )

    assert response.status_code in {200, 201}

    return response.json()["data"]


def _create_customer(
    client: TestClient,
    admin_headers: dict[str, str],
) -> dict:
    unique_value = uuid4().hex[:10].upper()

    response = client.post(
        "/api/v1/customers",
        headers=admin_headers,
        json={
            "customer_name": (
                f"Report Customer {unique_value}"
            ),
            "phone": "0812345678",
            "email": (
                f"report_{unique_value.lower()}"
                "@example.com"
            ),
            "address": "Bangkok",
        },
    )

    assert response.status_code in {200, 201}

    return response.json()["data"]


def _create_batch(
    client: TestClient,
    admin_headers: dict[str, str],
    product_id: int,
    *,
    quantity: int,
    expiry_days: int,
) -> dict:
    manufacturing_date = date.today()

    response = client.post(
        "/api/v1/batches",
        headers=admin_headers,
        json={
            "product_id": product_id,
            "lot_no": (
                f"REPORT-LOT-"
                f"{uuid4().hex[:12].upper()}"
            ),
            "mfg_date": (
                manufacturing_date.isoformat()
            ),
            "expiry_date": (
                manufacturing_date
                + timedelta(days=expiry_days)
            ).isoformat(),
            "quantity": quantity,
        },
    )

    assert response.status_code == 201

    return response.json()["data"]["batch"]


def _create_sales_order(
    client: TestClient,
    admin_headers: dict[str, str],
    customer_id: int,
    product_id: int,
) -> dict:
    response = client.post(
        "/api/v1/sales-orders/",
        headers=admin_headers,
        json={
            "customer_id": customer_id,
            "items": [
                {
                    "product_id": product_id,
                    "quantity": 2,
                    "unit_price": 150,
                }
            ],
        },
    )

    assert response.status_code == 200

    return response.json()["data"]


def test_stock_balance_report(
    warehouse_client: TestClient,
    admin_headers: dict[str, str],
) -> None:
    product = _create_product(
        warehouse_client,
        admin_headers,
        stock_qty=8,
        price=12.50,
    )

    response = warehouse_client.get(
        "/api/v1/reports/stock-balance"
    )

    assert response.status_code == 200

    data = response.json()

    item = next(
        (
            row
            for row in data
            if row["product_id"] == product["id"]
        ),
        None,
    )

    assert item is not None
    assert item["sku"] == product["sku"]
    assert item["stock_qty"] == 8
    assert float(item["price"]) == 12.50
    assert float(item["stock_value"]) == 100.00


def test_sales_summary_report(
    warehouse_client: TestClient,
    admin_headers: dict[str, str],
) -> None:
    before = warehouse_client.get(
        "/api/v1/reports/sales-summary"
    )

    assert before.status_code == 200

    before_data = before.json()

    customer = _create_customer(
        warehouse_client,
        admin_headers,
    )

    product = _create_product(
        warehouse_client,
        admin_headers,
    )

    _create_batch(
        warehouse_client,
        admin_headers,
        product["id"],
        quantity=10,
        expiry_days=90,
    )

    _create_sales_order(
        warehouse_client,
        admin_headers,
        customer["id"],
        product["id"],
    )

    after = warehouse_client.get(
        "/api/v1/reports/sales-summary"
    )

    assert after.status_code == 200

    after_data = after.json()

    assert after_data["total_orders"] == (
        before_data["total_orders"] + 1
    )

    assert after_data["total_sales_amount"] == (
        before_data["total_sales_amount"] + 300
    )


def test_stock_movement_report(
    warehouse_client: TestClient,
    admin_headers: dict[str, str],
) -> None:
    product = _create_product(
        warehouse_client,
        admin_headers,
    )

    stock_response = warehouse_client.post(
        "/api/v1/stock/in",
        headers=admin_headers,
        json={
            "product_id": product["id"],
            "quantity": 7,
            "remark": "Report movement test",
        },
    )

    assert stock_response.status_code == 200

    response = warehouse_client.get(
        "/api/v1/reports/stock-movement"
    )

    assert response.status_code == 200

    transactions = response.json()

    transaction = next(
        (
            row
            for row in transactions
            if (
                row["product_id"] == product["id"]
                and row["remark"]
                == "Report movement test"
            )
        ),
        None,
    )

    assert transaction is not None
    assert transaction["transaction_type"] == "IN"
    assert transaction["quantity"] == 7


def test_low_stock_report(
    warehouse_client: TestClient,
    admin_headers: dict[str, str],
) -> None:
    low_product = _create_product(
        warehouse_client,
        admin_headers,
        stock_qty=4,
    )

    high_product = _create_product(
        warehouse_client,
        admin_headers,
        stock_qty=20,
    )

    response = warehouse_client.get(
        "/api/v1/reports/low-stock?threshold=10"
    )

    assert response.status_code == 200

    data = response.json()

    product_ids = [
        item["product_id"]
        for item in data
    ]

    assert low_product["id"] in product_ids
    assert high_product["id"] not in product_ids

    for item in data:
        assert item["stock_qty"] <= 10
        assert item["threshold"] == 10


def test_expiring_report_excludes_expired(
    warehouse_client: TestClient,
    admin_headers: dict[str, str],
    db_session: Session,
) -> None:
    product = _create_product(
        warehouse_client,
        admin_headers,
    )

    expiring_batch = _create_batch(
        warehouse_client,
        admin_headers,
        product["id"],
        quantity=5,
        expiry_days=30,
    )

    future_batch = _create_batch(
        warehouse_client,
        admin_headers,
        product["id"],
        quantity=5,
        expiry_days=180,
    )

    expired_batch = ProductBatch(
        product_id=product["id"],
        lot_no=(
            f"REPORT-EXPIRED-"
            f"{uuid4().hex[:10].upper()}"
        ),
        mfg_date=(
            date.today() - timedelta(days=365)
        ),
        expiry_date=(
            date.today() - timedelta(days=1)
        ),
        quantity=5,
    )

    db_session.add(expired_batch)
    db_session.commit()
    db_session.refresh(expired_batch)

    response = warehouse_client.get(
        "/api/v1/reports/expiring?days=90"
    )

    assert response.status_code == 200

    data = response.json()

    batch_ids = [
        item["batch_id"]
        for item in data
    ]

    assert expiring_batch["id"] in batch_ids
    assert future_batch["id"] not in batch_ids
    assert expired_batch.id not in batch_ids


def test_sales_chart(
    warehouse_client: TestClient,
    admin_headers: dict[str, str],
) -> None:
    customer = _create_customer(
        warehouse_client,
        admin_headers,
    )

    product = _create_product(
        warehouse_client,
        admin_headers,
    )

    _create_batch(
        warehouse_client,
        admin_headers,
        product["id"],
        quantity=10,
        expiry_days=90,
    )

    _create_sales_order(
        warehouse_client,
        admin_headers,
        customer["id"],
        product["id"],
    )

    response = warehouse_client.get(
        "/api/v1/reports/chart/sales"
    )

    assert response.status_code == 200

    data = response.json()

    assert isinstance(data, list)
    assert len(data) >= 1

    for item in data:
        assert "date" in item
        assert "orders" in item
        assert "sales" in item


def test_stock_chart_sorted_descending(
    warehouse_client: TestClient,
    admin_headers: dict[str, str],
) -> None:
    _create_product(
        warehouse_client,
        admin_headers,
        stock_qty=5,
    )

    _create_product(
        warehouse_client,
        admin_headers,
        stock_qty=25,
    )

    response = warehouse_client.get(
        "/api/v1/reports/chart/stock"
    )

    assert response.status_code == 200

    data = response.json()

    quantities = [
        item["stock_qty"]
        for item in data
    ]

    assert quantities == sorted(
        quantities,
        reverse=True,
    )


def test_expiry_chart_sorted_ascending(
    warehouse_client: TestClient,
    admin_headers: dict[str, str],
) -> None:
    product = _create_product(
        warehouse_client,
        admin_headers,
    )

    _create_batch(
        warehouse_client,
        admin_headers,
        product["id"],
        quantity=5,
        expiry_days=60,
    )

    _create_batch(
        warehouse_client,
        admin_headers,
        product["id"],
        quantity=5,
        expiry_days=20,
    )

    response = warehouse_client.get(
        "/api/v1/reports/chart/expiry"
    )

    assert response.status_code == 200

    data = response.json()

    expiry_dates = [
        item["expiry_date"]
        for item in data
    ]

    assert expiry_dates == sorted(expiry_dates)


def _assert_excel_response(
    response,
    expected_filename: str,
) -> None:
    assert response.status_code == 200

    content_type = response.headers.get(
        "content-type",
        "",
    )

    assert (
        "spreadsheetml" in content_type
        or "application/octet-stream"
        in content_type
    )

    disposition = response.headers.get(
        "content-disposition",
        "",
    )

    assert expected_filename in disposition

    assert response.content.startswith(b"PK")


def test_export_stock_report(
    warehouse_client: TestClient,
    admin_headers: dict[str, str],
) -> None:
    _create_product(
        warehouse_client,
        admin_headers,
        stock_qty=5,
        price=25,
    )

    file_path = Path(
        "exports/stock_report.xlsx"
    )

    try:
        response = warehouse_client.get(
            "/api/v1/reports/export/stock"
        )

        _assert_excel_response(
            response,
            "stock_report.xlsx",
        )

        assert file_path.exists()

        workbook = load_workbook(file_path)
        worksheet = workbook["Stock Report"]

        assert worksheet.max_row >= 2
        assert worksheet["A1"].value == "Product ID"

        workbook.close()

    finally:
        if file_path.exists():
            file_path.unlink()


def test_export_sales_report(
    warehouse_client: TestClient,
    admin_headers: dict[str, str],
) -> None:
    customer = _create_customer(
        warehouse_client,
        admin_headers,
    )

    product = _create_product(
        warehouse_client,
        admin_headers,
    )

    _create_batch(
        warehouse_client,
        admin_headers,
        product["id"],
        quantity=10,
        expiry_days=90,
    )

    _create_sales_order(
        warehouse_client,
        admin_headers,
        customer["id"],
        product["id"],
    )

    file_path = Path(
        "exports/sales_report.xlsx"
    )

    try:
        response = warehouse_client.get(
            "/api/v1/reports/export/sales"
        )

        _assert_excel_response(
            response,
            "sales_report.xlsx",
        )

        workbook = load_workbook(file_path)
        worksheet = workbook["Sales Report"]

        assert worksheet.max_row >= 2
        assert worksheet["A1"].value == "SO Number"

        workbook.close()

    finally:
        if file_path.exists():
            file_path.unlink()


def test_export_low_stock_report(
    warehouse_client: TestClient,
    admin_headers: dict[str, str],
) -> None:
    _create_product(
        warehouse_client,
        admin_headers,
        stock_qty=3,
    )

    file_path = Path(
        "exports/low_stock_report.xlsx"
    )

    try:
        response = warehouse_client.get(
            (
                "/api/v1/reports/"
                "export/low-stock?threshold=10"
            )
        )

        _assert_excel_response(
            response,
            "low_stock_report.xlsx",
        )

        workbook = load_workbook(file_path)
        worksheet = workbook[
            "Low Stock Report"
        ]

        assert worksheet.max_row >= 2
        assert worksheet["E1"].value == "Threshold"

        workbook.close()

    finally:
        if file_path.exists():
            file_path.unlink()


def test_export_expiring_report(
    warehouse_client: TestClient,
    admin_headers: dict[str, str],
) -> None:
    product = _create_product(
        warehouse_client,
        admin_headers,
    )

    _create_batch(
        warehouse_client,
        admin_headers,
        product["id"],
        quantity=5,
        expiry_days=30,
    )

    file_path = Path(
        "exports/expiring_report.xlsx"
    )

    try:
        response = warehouse_client.get(
            (
                "/api/v1/reports/"
                "export/expiring?days=90"
            )
        )

        _assert_excel_response(
            response,
            "expiring_report.xlsx",
        )

        workbook = load_workbook(file_path)
        worksheet = workbook[
            "Expiring Report"
        ]

        assert worksheet.max_row >= 2
        assert worksheet["A1"].value == "Batch ID"
        assert worksheet["G1"].value == "Days Left"

        workbook.close()

    finally:
        if file_path.exists():
            file_path.unlink()