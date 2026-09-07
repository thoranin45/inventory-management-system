from app.core.dependencies import require_warehouse
from app.core.quantity import encode_quantities, quantity_text
from decimal import Decimal
from fastapi import APIRouter, Depends, Query
from sqlalchemy.orm import Session
from sqlalchemy import func

from app.core.batch_eligibility import business_today
from app.core.config import settings
from app.core.pagination import ListParams, list_params, phase8_json
from app.database import get_db
from datetime import date, datetime, timedelta
from app.models import Product, SalesOrder, StockTransaction, ProductBatch

from fastapi.responses import FileResponse
from openpyxl import Workbook
import os

router = APIRouter(dependencies=[Depends(require_warehouse)], prefix="/reports", tags=["Reports"])


def _report_page(items, total, params, message):
    import math
    return {
        "success": True,
        "message": message,
        "data": {
            "items": phase8_json(items),
            "pagination": {
                "page": params.page,
                "page_size": params.page_size,
                "total_items": total,
                "total_pages": math.ceil(total / params.page_size) if total else 0,
            },
        },
    }


@router.get("/stock-balance")
def stock_balance(db: Session = Depends(get_db)):
    products = db.query(Product).filter(
        Product.is_active == True
    ).all()

    return encode_quantities([
        {
            "product_id": p.id,
            "sku": p.sku,
            "product_name": p.product_name,
            "stock_qty": p.stock_qty,
            "price": float(p.price),
            "stock_value": float(p.price * p.stock_qty),
            "category_id": p.category_id,
        }
        for p in products
    ])


@router.get("/sales-summary")
def sales_summary(db: Session = Depends(get_db)):
    total_orders = db.query(SalesOrder).count()

    total_sales_amount = db.query(
        func.coalesce(func.sum(SalesOrder.total_amount), 0)
    ).scalar()

    return {
        "total_orders": total_orders,
        "total_sales_amount": float(total_sales_amount)
    }


@router.get("/stock-movement")
def stock_movement(
    db: Session = Depends(get_db),
    transaction_type: str | None = Query(default=None),
    date_from: datetime | None = Query(default=None),
    date_to: datetime | None = Query(default=None),
    params: ListParams = Depends(list_params),
):
    """Bounded StockTransaction history (Phase 8: no longer dumps the whole table)."""
    query = db.query(StockTransaction)
    if transaction_type:
        query = query.filter(StockTransaction.transaction_type == transaction_type)
    if date_from is not None:
        query = query.filter(StockTransaction.created_at >= date_from)
    if date_to is not None:
        query = query.filter(StockTransaction.created_at <= date_to)
    total = query.count()
    rows = (
        query.order_by(StockTransaction.created_at.desc(), StockTransaction.id.desc())
        .offset((params.page - 1) * params.page_size)
        .limit(params.page_size)
        .all()
    )
    return _report_page(
        [
            {
                "id": t.id,
                "product_id": t.product_id,
                "transaction_type": t.transaction_type,
                "quantity": t.quantity,
                "remark": t.remark,
                "created_at": getattr(t, "created_at", None),
            }
            for t in rows
        ],
        total,
        params,
        "Stock movement report",
    )


@router.get("/operational-stock")
def operational_stock_report(
    db: Session = Depends(get_db),
    params: ListParams = Depends(list_params),
):
    """Per-product owned vs operational-available vs expired vs near-expiry vs transit."""
    from app.repositories.product_repository import ProductRepository
    from app.repositories.stock_balance_repository import StockBalanceRepository
    from app.services.product_service import enrich_products, _PRODUCT_SORTS
    from app.core.pagination import paginate, resolve_ordering

    repo = ProductRepository(db)
    ordering = resolve_ordering(params, _PRODUCT_SORTS, "id", Product.id)
    items, total = paginate(repo.list_query(search=params.search), params, ordering)
    rows = enrich_products(
        StockBalanceRepository(db), items, business_today(), settings.near_expiry_days
    )
    return _report_page(rows, total, params, "Operational stock report")


@router.get("/low-stock-operational")
def low_stock_operational_report(db: Session = Depends(get_db)):
    from app.repositories.stock_balance_repository import StockBalanceRepository
    from app.services.attention_service import _low_stock_products, _low_stock_threshold

    rows = _low_stock_products(db, StockBalanceRepository(db), business_today())
    return encode_quantities({
        "items": [
            {
                "product_id": p.id,
                "sku": p.sku,
                "product_name": p.product_name,
                "operational_available_quantity": avail,
                "threshold": _low_stock_threshold(p),
            }
            for p, avail in rows
        ]
    })


def _batch_rows(batches, today):
    from app.core import batch_eligibility
    return [
        {
            "batch_id": b.id,
            "product_id": b.product_id,
            "lot_no": b.lot_no,
            "quantity": b.quantity,
            "expiry_date": b.expiry_date,
            "days_to_expiry": batch_eligibility.days_to_expiry(b, today),
        }
        for b in batches
    ]


@router.get("/expired-stock")
def expired_stock_report(db: Session = Depends(get_db)):
    from app.repositories.dashboard_repository import DashboardRepository

    today = business_today()
    return encode_quantities(
        {"items": _batch_rows(DashboardRepository(db).get_expired_batches(today), today)}
    )


@router.get("/near-expiry-stock")
def near_expiry_stock_report(
    db: Session = Depends(get_db),
    days: int = Query(default=None, ge=1, le=3650),
):
    from app.repositories.dashboard_repository import DashboardRepository

    today = business_today()
    window = today + timedelta(days=days or settings.near_expiry_days)
    batches = DashboardRepository(db).get_batches_expiring_between(today, window)
    return encode_quantities({"items": _batch_rows(batches, today)})


@router.get("/in-transit-stock")
def in_transit_stock_report(db: Session = Depends(get_db)):
    from app.repositories.stock_balance_repository import StockBalanceRepository

    rows = StockBalanceRepository(db).in_transit_query().all()
    return encode_quantities({
        "items": [
            {
                "id": r.id,
                "product_id": r.product_id,
                "batch_id": r.batch_id,
                "warehouse_id": r.warehouse_id,
                "location_id": r.location_id,
                "on_hand_qty": r.on_hand_qty,
            }
            for r in rows
        ]
    })

@router.get("/low-stock")
def low_stock_report(
    threshold: Decimal = Decimal("10.000"),
    db: Session = Depends(get_db)
):
    products = db.query(Product).filter(
        Product.is_active == True,
        Product.stock_qty <= threshold
    ).all()

    return encode_quantities([
        {
            "product_id": p.id,
            "sku": p.sku,
            "product_name": p.product_name,
            "stock_qty": p.stock_qty,
            "threshold": threshold
        }
        for p in products
    ])

@router.get("/expiring")
def expiring_report(
    days: int = 90,
    db: Session = Depends(get_db)
):
    today = business_today()
    target_date = today + timedelta(days=days)

    batches = db.query(ProductBatch).filter(
        ProductBatch.quantity > 0,
        ProductBatch.expiry_date >= today,
        ProductBatch.expiry_date <= target_date,
    ).order_by(
        ProductBatch.expiry_date.asc()
    ).all()

    return encode_quantities([
        {
            "batch_id": b.id,
            "product_id": b.product_id,
            "lot_no": b.lot_no,
            "quantity": b.quantity,
            "mfg_date": b.mfg_date,
            "expiry_date": b.expiry_date,
            "days_left": (b.expiry_date - today).days
        }
        for b in batches
    ])

@router.get("/chart/sales")
def sales_chart(db: Session = Depends(get_db)):
    data = db.query(
        func.date(SalesOrder.created_at).label("date"),
        func.count(SalesOrder.id).label("orders"),
        func.coalesce(func.sum(SalesOrder.total_amount), 0).label("sales")
    ).group_by(
        func.date(SalesOrder.created_at)
    ).order_by(
        func.date(SalesOrder.created_at)
    ).all()

    return encode_quantities([
        {
            "date": str(row.date),
            "orders": row.orders,
            "sales": float(row.sales)
        }
        for row in data
    ])


@router.get("/chart/stock")
def stock_chart(
    db: Session = Depends(get_db),
    limit: int = Query(default=20, ge=1, le=100),
):
    products = db.query(Product).filter(
        Product.is_active == True
    ).order_by(
        Product.stock_qty.desc(), Product.id.asc()
    ).limit(limit).all()

    return encode_quantities([
        {
            "product_name": p.product_name,
            "stock_qty": p.stock_qty
        }
        for p in products
    ])


@router.get("/chart/expiry")
def expiry_chart(
    db: Session = Depends(get_db),
    limit: int = Query(default=50, ge=1, le=200),
):
    from app.core import batch_eligibility

    today = business_today()
    near_cutoff = today + timedelta(days=settings.near_expiry_days)
    batches = db.query(ProductBatch).filter(
        ProductBatch.quantity > 0
    ).order_by(
        ProductBatch.expiry_date.asc().nulls_last(), ProductBatch.id.asc()
    ).limit(limit).all()

    def series(b):
        if b.expiry_date is None:
            return "no_expiry"
        if b.expiry_date < today:
            return "expired"
        if b.expiry_date <= near_cutoff:
            return "near_expiry"
        return "eligible"

    return encode_quantities([
        {
            "batch_id": b.id,
            "product_id": b.product_id,
            "lot_no": b.lot_no,
            "quantity": b.quantity,
            "expiry_date": str(b.expiry_date),
            "series": series(b),
        }
        for b in batches
    ])


@router.get("/sales")
def sales_report(
    db: Session = Depends(get_db),
    params: ListParams = Depends(list_params),
):
    from app.repositories.sales_order_repository import SalesOrderRepository
    from app.services.sales_order_service import get_sales_orders_service

    return get_sales_orders_service(SalesOrderRepository(db), params)


@router.get("/purchase-orders")
def purchase_orders_report(
    db: Session = Depends(get_db),
    params: ListParams = Depends(list_params),
):
    from app.repositories.purchase_order_repository import PurchaseOrderRepository
    from app.services.purchase_order_service import get_purchase_orders_service

    return get_purchase_orders_service(PurchaseOrderRepository(db), params)


@router.get("/transfers")
def transfers_report(
    db: Session = Depends(get_db),
    params: ListParams = Depends(list_params),
):
    from app.repositories.inventory_transfer_repository import InventoryTransferRepository
    from app.services.inventory_transfer_service import get_inventory_transfers_service

    return get_inventory_transfers_service(InventoryTransferRepository(db), params)

@router.get("/export/stock")
def export_stock_report(
    db: Session = Depends(get_db)
):

    products = db.query(Product).filter(
        Product.is_active == True
    ).all()

    os.makedirs(
        "exports",
        exist_ok=True
    )

    file_path = "exports/stock_report.xlsx"

    wb = Workbook()
    ws = wb.active

    ws.title = "Stock Report"

    ws.append([
        "Product ID",
        "SKU",
        "Product Name",
        "Stock Qty",
        "Price",
        "Stock Value"
    ])

    for p in products:

        ws.append([
            p.id,
            p.sku,
            p.product_name,
            quantity_text(p.stock_qty),
            float(p.price),
            float(p.price * p.stock_qty)
        ])

    wb.save(file_path)

    return FileResponse(
        file_path,
        filename="stock_report.xlsx"
    )

@router.get("/export/sales")
def export_sales_report(
    db: Session = Depends(get_db)
):

    orders = db.query(SalesOrder).all()

    os.makedirs(
        "exports",
        exist_ok=True
    )

    file_path = "exports/sales_report.xlsx"

    wb = Workbook()
    ws = wb.active

    ws.title = "Sales Report"

    ws.append([
        "SO Number",
        "Customer ID",
        "Status",
        "Total Amount",
        "Created At"
    ])

    for so in orders:

        ws.append([
            so.so_number,
            so.customer_id,
            so.status,
            float(so.total_amount),
            str(so.created_at)
        ])

    wb.save(file_path)

    return FileResponse(
        file_path,
        filename="sales_report.xlsx"
    )

@router.get("/export/low-stock")
def export_low_stock_report(
    threshold: Decimal = Decimal("10.000"),
    db: Session = Depends(get_db)
):
    products = db.query(Product).filter(
        Product.is_active == True,
        Product.stock_qty <= threshold
    ).all()

    os.makedirs("exports", exist_ok=True)
    file_path = "exports/low_stock_report.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "Low Stock Report"

    ws.append([
        "Product ID",
        "SKU",
        "Product Name",
        "Stock Qty",
        "Threshold"
    ])

    for p in products:
        ws.append([
            p.id,
            p.sku,
            p.product_name,
            quantity_text(p.stock_qty),
            format(threshold, "f")
        ])

    wb.save(file_path)

    return FileResponse(
        file_path,
        filename="low_stock_report.xlsx"
    )


@router.get("/export/expiring")
def export_expiring_report(
    days: int = 90,
    db: Session = Depends(get_db)
):
    today = business_today()
    target_date = today + timedelta(days=days)

    batches = db.query(ProductBatch).filter(
        ProductBatch.quantity > 0,
        ProductBatch.expiry_date >= today,
        ProductBatch.expiry_date <= target_date,
    ).order_by(
        ProductBatch.expiry_date.asc()
    ).all()

    os.makedirs("exports", exist_ok=True)
    file_path = "exports/expiring_report.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "Expiring Report"

    ws.append([
        "Batch ID",
        "Product ID",
        "Lot No",
        "Quantity",
        "MFG Date",
        "Expiry Date",
        "Days Left"
    ])

    for b in batches:
        ws.append([
            b.id,
            b.product_id,
            b.lot_no,
            quantity_text(b.quantity),
            str(b.mfg_date),
            str(b.expiry_date),
            (b.expiry_date - today).days
        ])

    wb.save(file_path)

    return FileResponse(
        file_path,
        filename="expiring_report.xlsx"
    )
